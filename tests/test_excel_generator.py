"""Tests for ExcelGenerator (9-sheet workbook output)."""

import openpyxl
from output.excel_generator import ExcelGenerator

EXPECTED_9_SHEETS = [
    "Executive Summary",
    "Tables & Columns",
    "Measures & Calculated Columns",
    "Relationships & Lineage",
    "Pages & Visuals",
    "Data Sources & Queries",
    "Usage & Unused Objects",
    "DAX Recommendations",
    "Model & Date Metadata",
]


def test_excel_generator_structure(tmp_path):
    output_file = str(tmp_path / "test_analysis.xlsx")

    metadata = {
        "model_info": {
            "Report / Project Name": "TestProject",
            "Semantic Model Name": "TestModel",
            "Model Format": "TMDL",
            "Culture": "en-US",
            "Compatibility Level": "1550",
            "Source Query Culture": "en-US",
            "Default Data Source Version": "powerBI_V3",
        },
        "tables": [{"Table Name": "DimCustomer", "Table Type": "Data", "Is Hidden": "No", "Data Category": "None", "Description": "None"}],
        "columns": [{"Table Name": "DimCustomer", "Column Name": "ID", "Data Type": "int64", "Is Hidden": "No", "Column Type": "Data", "Format String": "0", "Source Column": "ID", "Description": "None"}],
        "calculated_columns": [{"Table": "DimCustomer", "Column": "FullName", "Data Type": "string", "DAX Expression": "[First] & [Last]", "Is Hidden": "No"}],
        "measures": [{"Table": "DimCustomer", "Measure Name": "Count", "DAX Expression": "COUNTROWS(DimCustomer)", "Format String": "0", "Is Hidden": "No", "Description": "None"}],
        "relationships": [{"From Table": "Sales", "From Column": "ID", "To Table": "DimCustomer", "To Column": "ID", "Cardinality": "M:1", "Filter Direction": "Single", "Status": "Active"}],
        "key_columns": [{"Table": "Sales", "Column": "ID", "Relationship Role": "From (Many)", "Related Table": "DimCustomer", "Related Column": "ID"}],
        "relationship_metrics": {
            "Total Relationships": 1,
            "Active Relationships": 1,
            "Inactive Relationships": 0,
        },
        "pages": [{"Page ID": "p1", "Page Name": "Summary"}],
        "visuals": [{"Page ID": "p1", "Page Name": "Summary", "Visual ID": "v1", "Visual Name": "Card 1", "Visual Type": "Card", "Visual Title": "KPI"}],
        "data_sources": [{"Table": "Sales", "Source Type": "SQL Server", "Source Details": "Server: localhost", "Power Query / M Expression": "let ..."}],
        "date_table_summary": {"Date Table Exists": "No", "Date Table Name": "None"},
        "date_tables": [],
        "object_usage": [],
        "visual_usage": [],
        "dax_dependencies": [],
        "usage_lineage": [],
        "unused_objects": [],
        "dax_improvements": [],
    }

    ExcelGenerator.generate_report(metadata, output_file)

    wb = openpyxl.load_workbook(output_file)

    assert len(wb.sheetnames) == 9
    assert wb.sheetnames == EXPECTED_9_SHEETS

    # Verify Executive Summary
    summary_ws = wb["Executive Summary"]
    assert summary_ws["A1"].value == "Power BI Project Analysis — Executive Summary"
