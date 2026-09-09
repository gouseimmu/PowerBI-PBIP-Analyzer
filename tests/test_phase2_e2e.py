"""End-to-end tests for Phase 2 (6-sheet Excel generation & Flask results)."""

import openpyxl
import pytest
from app import app
from output.excel_generator import ExcelGenerator
from tests.fixtures import create_mock_tmdl_pbip_zip, create_mock_bim_pbip_zip

EXPECTED_6_SHEETS = [
    "Summary",
    "Model Inventory",
    "Report Usage",
    "Relationships & Lineage",
    "Data Sources",
    "DAX Analysis",
]


@pytest.fixture
def client():
    app.config["TESTING"] = True
    with app.test_client() as client:
        yield client


def test_excel_6_sheets_generation(tmp_path):
    output_file = str(tmp_path / "full_phase2_analysis.xlsx")

    metadata = {
        "project_name": "SalesAnalytics",
        "model_format": "TMDL",
        "report_format": "PBIR",
        "model_info": {
            "Report / Project Name": "SalesAnalytics",
            "Semantic Model Name": "SalesModel",
            "Model Format": "TMDL",
            "Culture": "en-US",
            "Compatibility Level": "1550",
            "Source Query Culture": "en-US",
            "Default Data Source Version": "powerBI_V3",
        },
        "tables": [{"Table Name": "Sales", "Table Type": "Data"}],
        "columns": [{"Table Name": "Sales", "Column Name": "Amount", "Data Type": "decimal", "Is Hidden": "No", "Column Type": "Data", "Format String": "$#,##0.00", "Source Column": "Amount", "Description": "None"}],
        "calculated_columns": [{"Table": "Sales", "Column": "Margin", "Data Type": "double", "DAX Expression": "Sales[Amount] * 0.2", "Is Hidden": "No"}],
        "measures": [{"Table": "Sales", "Measure Name": "Total Sales", "DAX Expression": "SUM(Sales[Amount])", "Format String": "$#,##0.00", "Is Hidden": "No", "Description": "None"}],
        "relationships": [{"From Table": "Sales", "From Column": "CustID", "To Table": "Customer", "To Column": "CustID", "Cardinality": "M:1", "Filter Direction": "Single", "Status": "Active"}],
        "key_columns": [{"Table": "Sales", "Column": "CustID", "Relationship Role": "From (Many)", "Related Table": "Customer", "Related Column": "CustID"}],
        "relationship_metrics": {
            "Total Relationships": 1,
            "Active Relationships": 1,
            "Inactive Relationships": 0,
        },
        "pages": [{"Page ID": "p1", "Page Name": "Overview"}],
        "visuals": [{"Page ID": "p1", "Page Name": "Overview", "Visual ID": "v1", "Visual Name": "Card 1", "Visual Type": "Card", "Visual Title": "Sales KPI"}],
        "data_sources": [{"Table": "Sales", "Source Type": "SQL Server", "Source Details": "Server: localhost", "Power Query / M Expression": "let ..."}],
        "date_table_summary": {"Date Table Exists": "No", "Date Table Name": "None"},
        "date_tables": [],
        "object_usage": [
            {"Object Type": "Measure", "Table": "Sales", "Object Name": "Total Sales", "Used": "Yes", "Usage Type": "Direct", "Direct Usage Count": 1, "Indirect Usage Count": 0, "Usage Location Count": 1}
        ],
        "visual_usage": [
            {"Page ID": "p1", "Page Name": "Overview", "Visual ID": "v1", "Visual Name": "Card 1", "Visual Type": "Card", "Visual Title": "Sales KPI", "Object Type": "Measure", "Table": "Sales", "Object Name": "Total Sales", "Usage Type": "Direct"}
        ],
        "dax_dependencies": [
            {"Source Object Type": "Measure", "Source Table": "Sales", "Source Object": "Total Sales", "Referenced Object Type": "Column", "Referenced Table": "Sales", "Referenced Object": "Amount", "Dependency Type": "Measure -> Column", "Dependency Level": 1, "DAX Expression": "SUM(Sales[Amount])", "Resolution Status": "Resolved"}
        ],
        "usage_lineage": [
            {"Page Name": "Overview", "Visual Name": "Card 1", "Visual Title": "Sales KPI", "Direct Object": "Total Sales", "Indirect Object": "Amount", "Dependency Path": "Total Sales -> Amount", "Dependency Level": 1, "DAX Expression": "SUM(Sales[Amount])"}
        ],
        "unused_objects": [
            {"Object Type": "Calculated Column", "Table": "Sales", "Object Name": "Margin", "Reason": "No direct or indirect report usage detected", "Direct Usage": "No", "Indirect Usage": "No"}
        ],
        "phase2_summary": {
            "Total Objects": 3,
            "Used Objects": 2,
            "Unused Objects": 1,
            "Directly Used Objects": 1,
            "Indirectly Used Objects": 1,
            "Direct Usage Records": 1,
            "Indirect Usage Records": 1,
            "Dependency Edges": 1,
            "Unresolved DAX References": 0,
            "Circular Dependencies": 0,
            "Visuals with Parsed References": 1,
            "Visuals with Unresolved References": 0,
        },
        "model_health": {"score": 95, "grade": "A", "status_text": "Good", "deductions": []},
    }

    ExcelGenerator.generate_report(metadata, output_file)

    wb = openpyxl.load_workbook(output_file)
    assert len(wb.sheetnames) == 6
    assert wb.sheetnames == EXPECTED_6_SHEETS


def test_flask_upload_phase2_full_pipeline(client, tmp_path):
    zip_path = str(tmp_path / "mock_tmdl_full.zip")
    create_mock_tmdl_pbip_zip(zip_path)

    with open(zip_path, "rb") as f:
        data = {"file": (f, "SalesAnalytics.zip")}
        response = client.post("/upload", data=data, content_type="multipart/form-data")

    assert response.status_code == 200
    assert b"SalesAnalytics" in response.data
    assert b"Extraction Successful" in response.data
    assert b"Download Excel Report" in response.data
