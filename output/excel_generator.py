"""Compact Excel report generator for the Power BI PBIP Analyzer.

Exactly six sheets are produced:
1. Summary
2. Model Inventory
3. Report Usage
4. Relationships & Lineage
5. Data Sources
6. DAX Analysis
"""
from __future__ import annotations

import os
import re
from typing import Any, Dict, List

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F497D")
HEADER_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color="1F497D")
SUBTITLE_FONT = Font(name="Segoe UI", size=10, italic=True, color="595959")
REGULAR_FONT = Font(name="Segoe UI", size=10)
BOLD_FONT = Font(name="Segoe UI", size=10, bold=True)
ZEBRA_FILL = PatternFill("solid", fgColor="F2F4F7")
SECTION_FILL = PatternFill("solid", fgColor="DCE6F1")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


class ExcelGenerator:
    @classmethod
    def generate_report(cls, metadata: Dict[str, Any], output_filepath: str) -> str:
        wb = Workbook()
        wb.remove(wb.active)
        cls._summary(wb, metadata)
        cls._model_inventory(wb, metadata)
        cls._report_usage(wb, metadata)
        cls._relationships(wb, metadata)
        cls._sources(wb, metadata)
        cls._dax(wb, metadata)
        os.makedirs(os.path.dirname(os.path.abspath(output_filepath)), exist_ok=True)
        wb.save(output_filepath)
        return output_filepath

    @classmethod
    def _summary(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        ws = wb.create_sheet("Summary")
        ws["A1"] = "Power BI PBIP Analyzer — Summary"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = "Compact model, report usage, lineage and DAX analysis"
        ws["A2"].font = SUBTITLE_FONT
        p2 = m.get("phase2_summary", {})
        p3 = m.get("phase3_summary", {})
        rel = m.get("relationship_metrics", {})
        rows = [
            ("Project Name", m.get("project_name", "Unknown")),
            ("Model Format", m.get("model_format", "Unknown")),
            ("Report Format", m.get("report_format", "Unknown")),
            ("Tables", len(m.get("tables", []))),
            ("Columns", len(m.get("columns", []))),
            ("Measures", len(m.get("measures", []))),
            ("Calculated Columns", len(m.get("calculated_columns", []))),
            ("Relationships", len(m.get("relationships", []))),
            ("Active Relationships", rel.get("Active Relationships", 0)),
            ("Pages", len(m.get("pages", []))),
            ("Visuals", len(m.get("visuals", []))),
            ("Used Objects", p2.get("Used Objects", 0)),
            ("Unused Objects", p2.get("Unused Objects", 0)),
            ("DAX Dependencies", p2.get("Dependency Edges", 0)),
            ("Unresolved DAX References", p2.get("Unresolved DAX References", 0)),
            ("Circular Dependencies", p2.get("Circular Dependencies", 0)),
            ("DAX Objects Analyzed", p3.get("DAX Objects Analyzed", 0)),
            ("Optimization Opportunities", p3.get("Objects With Potential Improvements", 0)),
            ("AI Provider", m.get("ai_provider", "Unknown")),
        ]
        cls._write_kv(ws, 4, rows)
        findings = []
        if p2.get("Unresolved DAX References", 0):
            findings.append(("DAX", "Unresolved DAX references require review."))
        if p2.get("Circular Dependencies", 0):
            findings.append(("DAX", "Circular dependencies detected."))
        if p2.get("Unused Objects", 0):
            findings.append(("Usage", f"{p2.get('Unused Objects')} objects have no direct or indirect report usage."))
        if not findings:
            findings.append(("Status", "No major findings detected by deterministic checks."))
        start = 4 + len(rows) + 2
        ws.cell(start, 1, "Important Findings").fill = SECTION_FILL
        ws.cell(start, 1).font = BOLD_FONT
        start += 1
        for cat, text in findings:
            ws.cell(start, 1, cat).font = BOLD_FONT
            ws.cell(start, 2, text).font = REGULAR_FONT
            start += 1
        ws.column_dimensions["A"].width = 34
        ws.column_dimensions["B"].width = 85

    @classmethod
    def _model_inventory(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        usage = {(u.get("Object Type"), u.get("Table"), u.get("Object Name")): u for u in m.get("object_usage", [])}
        locations: Dict[tuple, Dict[str, set]] = {}
        for v in m.get("visual_usage", []):
            key = (v.get("Object Type"), v.get("Table"), v.get("Object Name"))
            locations.setdefault(key, {"pages": set(), "visuals": set()})
            if v.get("Page Name"): locations[key]["pages"].add(v["Page Name"])
            if v.get("Visual Title") or v.get("Visual Name"):
                locations[key]["visuals"].add(v.get("Visual Title") or v.get("Visual Name"))
        calc_map = {(x.get("Table"), x.get("Column")): x for x in m.get("calculated_columns", [])}
        rows = []
        # Tables: used if any child semantic object is used.
        child_used = {}
        for u in m.get("object_usage", []):
            if u.get("Used") == "Yes": child_used[u.get("Table")] = True
        for t in m.get("tables", []):
            name = t.get("Table Name") or t.get("name") or t.get("Table", "")
            rows.append({"Table": name, "Object": name, "Object Type": "Table", "Data Type": "", "DAX Expression": "", "Hidden": t.get("Is Hidden", "No"), "Used": "Yes" if child_used.get(name) else "No", "Usage Type": "Used via child objects" if child_used.get(name) else "Unused", "Pages": "", "Visuals": ""})
        for c in m.get("columns", []):
            table, obj = c.get("Table Name", ""), c.get("Column Name", "")
            is_calc = c.get("Column Type") == "Calculated" or (table, obj) in calc_map
            typ = "Calculated Column" if is_calc else "Column"
            u = usage.get((typ, table, obj), {})
            loc = locations.get((typ, table, obj), {"pages": set(), "visuals": set()})
            rows.append({"Table": table, "Object": obj, "Object Type": typ, "Data Type": c.get("Data Type", ""), "DAX Expression": calc_map.get((table, obj), {}).get("DAX Expression", "") if is_calc else "", "Hidden": c.get("Is Hidden", "No"), "Used": u.get("Used", "No"), "Usage Type": u.get("Usage Type", "Unused"), "Pages": "; ".join(sorted(loc["pages"])), "Visuals": "; ".join(sorted(loc["visuals"]))})
        for x in m.get("measures", []):
            table, obj = x.get("Table", ""), x.get("Measure Name", "")
            u = usage.get(("Measure", table, obj), {})
            loc = locations.get(("Measure", table, obj), {"pages": set(), "visuals": set()})
            rows.append({"Table": table, "Object": obj, "Object Type": "Measure", "Data Type": "", "DAX Expression": x.get("DAX Expression", ""), "Hidden": x.get("Is Hidden", "No"), "Used": u.get("Used", "No"), "Usage Type": u.get("Usage Type", "Unused"), "Pages": "; ".join(sorted(loc["pages"])), "Visuals": "; ".join(sorted(loc["visuals"]))})
        headers = ["Table", "Object", "Object Type", "Data Type", "DAX Expression", "Hidden", "Used", "Usage Type", "Pages", "Visuals"]
        cls._table(wb, "Model Inventory", "Semantic Model Inventory & Report Usage", rows, headers)

    @classmethod
    def _report_usage(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        rows = []
        for v in m.get("visual_usage", []):
            rows.append({
                "Page": v.get("Page Name", ""), "Visual ID": v.get("Visual ID", ""),
                "Visual": v.get("Visual Name", ""), "Visual Type": v.get("Visual Type", ""),
                "Visual Title": v.get("Visual Title", ""), "Object Type": v.get("Object Type", ""),
                "Table": v.get("Table", ""), "Object": v.get("Object Name", ""),
                "Usage Type": v.get("Usage Type", "Direct")
            })
        headers = ["Page", "Visual ID", "Visual", "Visual Type", "Visual Title", "Object Type", "Table", "Object", "Usage Type"]
        cls._table(wb, "Report Usage", "Report Visual → Semantic Object Usage", rows, headers)

    @classmethod
    def _relationships(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        rows = []
        for r in m.get("relationships", []):
            rows.append({"Type": "Relationship", "From": f"{r.get('From Table','')}[{r.get('From Column','')}]", "To": f"{r.get('To Table','')}[{r.get('To Column','')}]", "Cardinality": r.get("Cardinality", ""), "Filter Direction": r.get("Filter Direction", ""), "Status": r.get("Status", "")})
        for d in m.get("dax_dependencies", []):
            rows.append({"Type": "DAX Dependency", "From": f"{d.get('Source Table','')}[{d.get('Source Object','')}]", "To": f"{d.get('Referenced Table','')}[{d.get('Referenced Object','')}]", "Cardinality": "", "Filter Direction": "", "Status": d.get("Resolution Status", "Resolved")})
        headers = ["Type", "From", "To", "Cardinality", "Filter Direction", "Status"]
        cls._table(wb, "Relationships & Lineage", "Relationships and DAX Dependency Lineage", rows, headers)

    @classmethod
    def _sources(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        rows = []
        for s in m.get("data_sources", []):
            details = s.get("Source Details", "")
            rows.append({"Table": s.get("Table", ""), "Source Type": s.get("Source Type", "Other"), "Source Details": details, "Power Query / M Expression": s.get("Power Query / M Expression", "")})
        headers = ["Table", "Source Type", "Source Details", "Power Query / M Expression"]
        cls._table(wb, "Data Sources", "Data Sources & Power Query", rows, headers)

    @classmethod
    def _dax(cls, wb: Workbook, m: Dict[str, Any]) -> None:
        analysis = m.get("dax_analysis", [])
        improvements = m.get("dax_improvements", [])
        by_key = {(x.get("Table"), x.get("Object Name")): x for x in improvements}
        rows = []
        for x in analysis:
            key = (x.get("Table"), x.get("Object Name"))
            imp = by_key.get(key, {})
            rows.append({"Table": x.get("Table", ""), "Object": x.get("Object Name", ""), "Object Type": x.get("Object Type", ""), "DAX Expression": x.get("DAX Expression", x.get("Original DAX", "")), "Complexity Score": x.get("Complexity Score", ""), "Issue": imp.get("Issue", x.get("Issue", "")), "Severity": imp.get("Severity", x.get("Severity", "")), "Recommendation": imp.get("Recommendation", x.get("Recommendation", "")), "Suggested DAX": imp.get("Suggested DAX", x.get("Suggested DAX", "")), "Manual Review": "Yes" if imp.get("Requires Manual Validation") or imp.get("Validation Status") == "Needs Review" else "No"})
        if not rows:
            for x in improvements:
                rows.append({"Table": x.get("Table", ""), "Object": x.get("Object Name", ""), "Object Type": x.get("Object Type", ""), "DAX Expression": x.get("Original DAX", ""), "Complexity Score": x.get("Complexity Score", ""), "Issue": x.get("Issue", ""), "Severity": x.get("Severity", ""), "Recommendation": x.get("Recommendation", ""), "Suggested DAX": x.get("Suggested DAX", ""), "Manual Review": "Yes" if x.get("Requires Manual Validation") or x.get("Validation Status") == "Needs Review" else "No"})
        headers = ["Table", "Object", "Object Type", "DAX Expression", "Complexity Score", "Issue", "Severity", "Recommendation", "Suggested DAX", "Manual Review"]
        cls._table(wb, "DAX Analysis", "DAX Quality, Complexity & Recommendations", rows, headers)

    @staticmethod
    def _write_kv(ws, start: int, rows: List[tuple]) -> None:
        for i, (k, v) in enumerate(rows, start=start):
            ws.cell(i, 1, k).font = BOLD_FONT
            ws.cell(i, 2, str(v)).font = REGULAR_FONT
            ws.cell(i, 1).border = ws.cell(i, 2).border = BORDER
            if (i - start) % 2:
                ws.cell(i, 1).fill = ws.cell(i, 2).fill = ZEBRA_FILL

    @classmethod
    def _table(cls, wb: Workbook, name: str, title: str, rows: List[Dict[str, Any]], headers: List[str]) -> None:
        ws = wb.create_sheet(name)
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = f"Total Items: {len(rows)}"
        ws["A2"].font = SUBTITLE_FONT
        for j, h in enumerate(headers, 1):
            c = ws.cell(4, j, h)
            c.fill = HEADER_FILL
            c.font = HEADER_FONT
            c.alignment = Alignment(wrap_text=True, vertical="center")
            c.border = BORDER
        for i, row in enumerate(rows, 5):
            for j, h in enumerate(headers, 1):
                c = ws.cell(i, j, row.get(h, "") if row.get(h, "") not in (None, "") else "")
                c.font = REGULAR_FONT
                c.border = BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
                if (i - 5) % 2:
                    c.fill = ZEBRA_FILL
        ws.freeze_panes = "A5"
        end = max(4, 4 + len(rows))
        ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{end}"
        for j, h in enumerate(headers, 1):
            max_len = len(h)
            for i in range(4, end + 1):
                val = str(ws.cell(i, j).value or "")
                max_len = max(max_len, min(len(val.split("\n")[0]), 60))
            ws.column_dimensions[get_column_letter(j)].width = min(max(max_len + 3, 14), 55)
